import re
from typing import Match
from langdetect import detect, LangDetectException
from fpdf import FPDF
from docx import Document


def detectar_idioma(texto: str) -> str:
    try:
        return detect(texto)
    except LangDetectException:
        return "unknown"

def destacar_texto(texto: str, gravidade: dict, equivalencias: dict):
    abusivas = []
    frases_encontradas = set()
    texto_destacado = texto

    for palavra in gravidade['grave'] + gravidade['moderado']:
        padrao = re.compile(palavra, re.IGNORECASE)
        for m in padrao.finditer(texto):
            nivel = 'grave' if palavra in gravidade['grave'] else 'moderado'
            inicio = texto.rfind('.', 0, m.start()) + 1
            fim = texto.find('.', m.end())
            if fim == -1: fim = len(texto)
            frase = texto[inicio:fim].strip()

            if frase not in frases_encontradas:
                frases_encontradas.add(frase)
                abusivas.append({'clausula': frase, 'nivel': nivel})

            def destacar(match: Match[str]) -> str:
                return f'<span style="color:red;font-weight:bold;">{match.group(0)}</span>'

            texto_destacado = re.sub(padrao, destacar, texto_destacado)

    for clausula, frases in equivalencias.items():
        for frase in frases:
            if frase.lower() in texto.lower():
                abusivas.append({'clausula': frase, 'nivel': 'grave'})
                texto_destacado = texto_destacado.replace(
                    frase,
                    f'<span style="color:red;font-weight:bold;">{frase}</span>'
                )

    return texto_destacado, abusivas

def gerar_pdf(texto_destacado, abusivas, resumo):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    texto_limpo = re.sub(r'<.*?>', '', texto_destacado).replace('€', 'EUR')
    pdf.multi_cell(0, 10, "Texto analisado:\n" + texto_limpo)
    pdf.ln(5)

    pdf.cell(0, 10, "Clausulas encontradas:", ln=True)
    for clausula in abusivas:
        nivel = "Grave" if clausula['nivel'] == 'grave' else "Moderado"
        clausula_limpa = clausula['clausula'].replace('€', 'EUR')
        pdf.multi_cell(0, 10, f"- {clausula_limpa} ({nivel})")

    pdf.ln(5)
    pdf.cell(0, 10, "Resumo:", ln=True)
    pdf.multi_cell(0, 10, resumo.replace('€', 'EUR'))

    return pdf.output(dest='S').encode('latin1')

def gerar_docx(texto_destacado, abusivas, resumo):
    doc = Document()
    doc.add_heading("Relatório de Análise de Cláusulas Abusivas", level=1)

    doc.add_heading("Texto analisado:", level=2)
    doc.add_paragraph(re.sub(r'<.*?>', '', texto_destacado))

    doc.add_heading("Cláusulas encontradas:", level=2)
    for clausula in abusivas:
        nivel = "Grave" if clausula['nivel'] == 'grave' else "Moderado"
        doc.add_paragraph(f"- {clausula['clausula']} ({nivel})")

    doc.add_heading("Resumo:", level=2)
    doc.add_paragraph(resumo)

    from io import BytesIO
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
